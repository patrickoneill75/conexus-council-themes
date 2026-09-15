"""Reading a survey export's rows regardless of whether it's a .csv or .xlsx file.

Both survey_extract.py (natural-language extraction) and quant_extract.py (the quant
unpivot) read from the same New Survey Directory now, and real exports from the survey
tool come out as .csv while some historical files are .xlsx -- this is the one place
that decides which reader to use, so neither module has to know or care.
"""
from __future__ import annotations

import csv
import io

from openpyxl import load_workbook


def read_rows(filename: str, content: bytes) -> tuple[list[str], list[list]]:
    """-> (header, data_rows). Every data row is padded to len(header) so callers can
    index positionally the same way regardless of source format."""
    if filename.lower().endswith(".csv"):
        text = content.decode("utf-8-sig", errors="replace")
        rows = list(csv.reader(io.StringIO(text)))
        if not rows:
            return [], []
        header = [cell.strip() for cell in rows[0]]
        data = [_pad(row, len(header)) for row in rows[1:]]
        return header, data

    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = [str(cell).strip() if cell is not None else "" for cell in next(rows_iter)]
    except StopIteration:
        return [], []
    return header, [_pad(list(row), len(header)) for row in rows_iter]


def _pad(row: list, width: int) -> list:
    if len(row) < width:
        return row + [None] * (width - len(row))
    return row
