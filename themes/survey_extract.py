"""Reading a raw post-meeting survey export, ready for Claude to pull feedback items
out of.

What counts as a "free response" question versus a rating question is left entirely
to Claude at the extraction step — not decided here by column name, since survey
templates change wording and question order between quarters. This module's only job
is to strip out pure metadata (timestamps, IP address, browser, respondent identity)
that has no business reaching the API and would just be noise in the prompt.
"""
from __future__ import annotations

import io
import re

from openpyxl import load_workbook

_DROP = re.compile(
    r"^time$|ip address|unique id|\blocation\b|\bbrowser\b|name \(first\)|name \(last\)",
    re.IGNORECASE,
)


def response_rows(survey_bytes: bytes) -> list[dict]:
    """Every response, as {question: answer}, blank answers and metadata columns
    dropped. The first row is taken as the header regardless of sheet name."""
    wb = load_workbook(io.BytesIO(survey_bytes), data_only=True)
    ws = wb.active
    header = [str(c.value).strip() if c.value is not None else "" for c in next(ws.iter_rows(max_row=1))]
    keep = [i for i, h in enumerate(header) if h and not _DROP.search(h)]

    rows = []
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not any(raw):
            continue
        row = {header[i]: raw[i] for i in keep if raw[i] not in (None, "")}
        if row:
            rows.append(row)
    return rows
