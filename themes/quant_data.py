"""Reading the three fixed-name files a dashboard run finds in the Data Folder:
'Council Meeting Helper.csv', 'Post-Meeting Survey.csv', and 'Content
Categories.xlsx'.

The Helper and Survey exports are cumulative — every fresh export from the survey
tool contains every meeting/response ever collected, not just the newest — so each
upload replaces the file in Box wholesale rather than merging. Content Categories is
static/slow-changing and admin-maintained directly in Box.
"""
from __future__ import annotations

import csv
import io
from datetime import date

from openpyxl import load_workbook

from . import quant_extract

HELPER_FILENAME = "Council Meeting Helper.csv"
SURVEY_FILENAME = "Post-Meeting Survey.csv"
CATEGORIES_FILENAME = "Content Categories.xlsx"

# quant_extract.as_date() is the one canonical date parser both this module and
# quant_extract's own unpivot() use -- keeping a second, separately-maintained format
# list here is exactly how the Helper file's real "12-Aug-26" (%d-%b-%y) style dates
# went unrecognized even though quant_extract already handled that format.
parse_date = quant_extract.as_date


def _int_or_none(value: str) -> int | None:
    value = (value or "").strip()
    if not value:
        return None
    try:
        return int(float(value))
    except ValueError:
        return None


def read_helper(content: bytes) -> dict[date, dict]:
    """-> { Meeting Date: {year, quarter, region, total_registrants, total_attendees} }.

    A meeting with no attendance numbers recorded yet still gets an entry (Year/Quarter/
    Region are what a survey row needs to resolve; the attendee counts are optional).
    """
    text = content.decode("utf-8-sig", errors="replace")
    reader = csv.DictReader(io.StringIO(text))
    fieldnames = {(f or "").strip() for f in (reader.fieldnames or [])}
    date_col = next((f for f in fieldnames if f.lower() == "meeting date"), None)
    if date_col is None:
        # Every row's Meeting Date lookup below depends on this exact column existing --
        # its absence alone means the whole file resolves to zero meetings, which (since
        # the quant dashboard is rebuilt from scratch every run -- see quant_publish.py)
        # would silently wipe out everything already published. Almost always means the
        # uploaded file isn't really a Helper export (wrong file picked, a spreadsheet
        # saved in some other format and renamed to .csv on the way in, or a real export
        # whose column got renamed) -- fail loudly instead of silently returning {}.
        raise ValueError(
            f"'{HELPER_FILENAME}' has no 'Meeting Date' column -- nothing in it can "
            f"resolve. Its header reads: {sorted(fieldnames)!r}. Check the uploaded file "
            "is really a .csv export from the Helper form."
        )

    out: dict[date, dict] = {}
    unparseable: list[str] = []
    for row in reader:
        raw_date = row.get(date_col, "")
        meeting_date = parse_date(raw_date)
        if not meeting_date:
            if (raw_date or "").strip():
                unparseable.append(raw_date)
            continue
        out[meeting_date] = {
            "year": _int_or_none(row.get("Year", "")),
            "quarter": (row.get("Quarter") or "").strip(),
            "region": (row.get("Region") or "").strip(),
            "total_registrants": _int_or_none(row.get("Total Registrants", "")),
            "total_attendees": _int_or_none(row.get("Total Attendees", "")),
        }

    if not out and unparseable:
        # The column exists, and rows have real values in it, but none of them matched
        # any known format -- almost certainly a date format this hasn't seen before
        # (parse_date()'s format list is a fixed, known set). Surfacing real examples
        # here is what makes that fixable instead of just "0 meetings" with no clue why.
        raise ValueError(
            f"'{HELPER_FILENAME}' has a '{date_col}' column, but none of its values "
            f"parsed as a date -- e.g. {unparseable[:3]!r}. Recognized formats: "
            f"{quant_extract._DATE_FORMATS!r}."
        )
    return out


def read_categories(content: bytes) -> dict[str, str]:
    """-> { Metric name ("Detailed Category" in the source file): Content Category }.

    A metric with no mapping here (the source file doesn't cover every possible
    "Panel N"/"Workshop N" number) is simply left out of byCategory downstream —
    matching the real model's relationship exactly rather than inventing a group.
    """
    wb = load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    header = [str(c.value).strip() if c.value is not None else ""
              for c in next(ws.iter_rows(max_row=1))]
    detail_col = header.index("Detailed Category")
    content_col = header.index("Content Category")
    out: dict[str, str] = {}
    for raw in ws.iter_rows(min_row=2, values_only=True):
        if not any(raw):
            continue
        detail, content_cat = raw[detail_col], raw[content_col]
        if detail and content_cat:
            out[str(detail).strip()] = str(content_cat).strip()
    return out
