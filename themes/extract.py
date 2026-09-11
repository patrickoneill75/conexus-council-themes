"""Turning a Word document into the structure the dashboard renders.

The output is plain text in a small tree — headings with blocks beneath them — rather
than HTML. The page builds DOM nodes and sets textContent, so nothing anyone types into
a meeting document can ever be interpreted as markup.
"""
from __future__ import annotations

import io
import re

from docx import Document


def list_level(paragraph) -> int | None:
    """Word's own list level for a paragraph, or None if it isn't a list item.

    Word marks list items with a numbering property rather than a style name, which is
    why checking the style alone misses bullets in some templates — including Conexus's,
    if the documents were started from a house template.
    """
    p_pr = paragraph._p.pPr
    if p_pr is None or p_pr.numPr is None:
        if "list" in (paragraph.style.name or "").lower():
            return 0
        return None
    ilvl = p_pr.numPr.ilvl
    try:
        return int(ilvl.val) if ilvl is not None else 0
    except (TypeError, ValueError):
        return 0


def heading_level(paragraph) -> int | None:
    """1-9 for a Word heading, 0 for Title, None otherwise."""
    name = (paragraph.style.name or "").strip()
    if name.lower() == "title":
        return 0
    match = re.match(r"^Heading (\d)$", name, re.IGNORECASE)
    return int(match.group(1)) if match else None


def sections(docx_bytes: bytes) -> list[dict]:
    """Sections: a heading plus the blocks beneath it.

    Content before the first heading goes into an untitled opening section, so a lead
    paragraph — or a document written as one flat bullet list with no headings at all —
    is still published rather than silently dropped.
    """
    doc = Document(io.BytesIO(docx_bytes))
    out: list[dict] = []
    current = {"title": None, "blocks": []}

    for paragraph in doc.paragraphs:
        text = " ".join(paragraph.text.split())
        if not text:
            continue

        level = heading_level(paragraph)
        if level is not None:
            if current["title"] is not None or current["blocks"]:
                out.append(current)
            current = {"title": text, "level": max(level, 1), "blocks": []}
            continue

        depth = list_level(paragraph)
        if depth is None:
            current["blocks"].append({"type": "p", "text": text})
        else:
            current["blocks"].append({"type": "bullet", "depth": min(depth, 3), "text": text})

    if current["title"] is not None or current["blocks"]:
        out.append(current)

    # A table of notes is still content worth keeping.
    for table in doc.tables:
        rows = []
        for row in table.rows:
            cells = [" ".join(c.text.split()) for c in row.cells]
            line = " — ".join([c for c in cells if c])
            if line:
                rows.append({"type": "bullet", "depth": 0, "text": line})
        if rows:
            out.append({"title": "From the table", "level": 1, "blocks": rows})

    return out


def count_bullets(section_list: list[dict]) -> int:
    return sum(
        len([b for b in s["blocks"] if b["type"] == "bullet"]) for s in section_list
    )
